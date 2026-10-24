"""Plan semanal TIC.

Arranque local: python app.py
Luego abrir http://127.0.0.1:5050

En el servidor: python app.py --servidor
Abre http://192.168.20.205:5099
"""

from __future__ import annotations

import os
import re
import sqlite3
import sys
import threading
import time
import unicodedata
import webbrowser
from datetime import date, datetime, timedelta

from flask import Flask, g, jsonify, render_template, request

BASE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE, "data", "plan.db")
SCHEMA = os.path.join(BASE, "schema.sql")
EXCEL = os.path.join(BASE, "PLAN DE W SEMANAL TECNOLOGIA 2026.xlsx")
ANIO_EXCEL = 2026

ESTADOS = ("Sin estado", "Pendiente", "En curso", "Notificado", "Realizado")
SECCIONES = ("prioritaria", "extra")

CANON = {
    "brayan": "Brayan Gomez",
    "brayan gomez": "Brayan Gomez",
    "juan": "Juan Carreño",
    "juan carreno": "Juan Carreño",
    "juan cerreno": "Juan Carreño",
    "daniel": "Daniel Almeida",
    "daniel almeida": "Daniel Almeida",
    "leonardo": "Leonardo Lara",
    "leonardo lara": "Leonardo Lara",
    "johan": "Johan Pinto",
    "johan pinto": "Johan Pinto",
    "todos": "Todos",
    "desarrollo": "Desarrollo",
}

EQUIPO = (
    "Brayan Gomez",
    "Juan Carreño",
    "Daniel Almeida",
    "Leonardo Lara",
    "Johan Pinto",
    "Desarrollo",
    "Todos",
)

app = Flask(__name__)
app.config["TEMPLATES_AUTO_RELOAD"] = True
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0
app.json.ensure_ascii = False


def fold(value):
    text = unicodedata.normalize("NFD", value or "")
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", text).strip().lower()


def texto(value):
    if value is None:
        return None
    cleaned = str(value).strip()
    return cleaned or None


def connect():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def get_db():
    if "db" not in g:
        g.db = connect()
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def limites_semana(anio, numero):
    inicio = date.fromisocalendar(anio, numero, 1)
    fin = inicio + timedelta(days=6)
    return inicio.isoformat(), fin.isoformat()


def siguiente_semana(anio, numero):
    try:
        date.fromisocalendar(anio, numero + 1, 1)
        return anio, numero + 1
    except ValueError:
        return anio + 1, 1


def as_fecha(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def as_prioridad(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        numero = int(float(value))
    except (TypeError, ValueError):
        return None
    if 1 <= numero <= 53:
        return numero
    return None


def norm_estado(value):
    if value is None or str(value).strip() == "":
        return "Sin estado", None
    raw = str(value).strip()
    low = raw.lower()
    if low == "realizado":
        return "Realizado", None
    if low == "pendiente":
        return "Pendiente", None
    if low.startswith("notificado"):
        return "Notificado", raw
    return "Sin estado", raw


def titulo_foco(value, orden):
    raw = texto(value) or f"Actividad {orden}"
    return re.sub(r"(?i)\bactuvidad\b", "Actividad", raw)


def separar_personas(raw):
    if not raw or not str(raw).strip():
        return []
    text = re.sub(r"\s+y\s+", ",", str(raw).strip(), flags=re.I)
    parts = re.split(r"\s*(?:,|/|;|\||-)\s*", text)
    names = []
    for part in parts:
        part = re.sub(r"\s+", " ", part).strip(" .")
        if len(part) < 2:
            continue
        name = CANON.get(fold(part), part)
        if name not in names:
            names.append(name)
    return names


def persona_por_nombre(db, nombre):
    key = fold(nombre)
    for row in db.execute("SELECT * FROM personal"):
        if fold(row["nombre"]) == key:
            return row
    return None


def obtener_o_crear_persona(db, nombre):
    cleaned = re.sub(r"\s+", " ", (nombre or "")).strip()
    if len(cleaned) < 2:
        raise ValueError("Escribe el nombre de la persona.")
    if len(cleaned) > 80:
        raise ValueError("El nombre es demasiado largo.")
    if not re.fullmatch(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ][A-Za-zÁÉÍÓÚÜÑáéíóúüñ .'\-]{1,79}", cleaned):
        raise ValueError("El nombre solo puede tener letras, espacios y guiones.")
    final = CANON.get(fold(cleaned), cleaned)
    existente = persona_por_nombre(db, final)
    if existente:
        return existente
    cur = db.execute("INSERT INTO personal (nombre) VALUES (?)", (final,))
    return db.execute("SELECT * FROM personal WHERE id = ?", (cur.lastrowid,)).fetchone()


def importar_excel(conn):
    import openpyxl

    if not os.path.exists(EXCEL):
        raise FileNotFoundError(EXCEL)

    cache = {}

    def id_persona(nombre):
        key = fold(nombre)
        canon = CANON.get(key, nombre)
        ckey = fold(canon)
        if ckey in cache:
            return cache[ckey]
        cur = conn.execute("INSERT INTO personal (nombre) VALUES (?)", (canon,))
        cache[ckey] = cur.lastrowid
        cache[key] = cur.lastrowid
        return cur.lastrowid

    for nombre in EQUIPO:
        id_persona(nombre)

    wb = openpyxl.load_workbook(EXCEL, data_only=True, read_only=True)
    semanas = 0
    tareas = 0
    focos = 0
    try:
        for hoja in wb.sheetnames:
            match = re.search(r"(\d+)", hoja)
            if not match:
                continue
            numero = int(match.group(1))
            if not 1 <= numero <= 53:
                continue
            inicio, fin = limites_semana(ANIO_EXCEL, numero)
            es_copia = 1 if hoja.lower().startswith("copia") else 0
            try:
                cur = conn.execute(
                    """
                    INSERT INTO semanas (anio, numero, nombre_hoja, fecha_inicio, fecha_fin, es_copia)
                    VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (ANIO_EXCEL, numero, hoja, inicio, fin, es_copia),
                )
            except sqlite3.IntegrityError:
                continue
            semana_id = cur.lastrowid
            semanas += 1
            ws = wb[hoja]
            seccion = None
            orden = {"prioritaria": 0, "extra": 0}
            foco_n = 0
            for row in ws.iter_rows(min_col=1, max_col=11, values_only=True):
                raw = row[0]
                if raw is None:
                    continue
                text = str(raw).strip()
                if not text or text == "@":
                    continue
                upper = text.upper()
                if upper == "ACTIVIDADES":
                    continue
                if upper == "ACTIVIDADES PRIORITARIAS":
                    seccion = "prioritaria"
                    continue
                if upper == "ACTIVIDADES EXTRAS":
                    seccion = "extra"
                    continue
                if seccion is None and upper.startswith("ACT"):
                    foco_n += 1
                    descripcion = texto(row[1])
                    if not descripcion:
                        continue
                    conn.execute(
                        """
                        INSERT INTO focos (semana_id, orden, titulo, descripcion, estado)
                        VALUES (?, ?, ?, ?, 'Sin estado')
                        """,
                        (semana_id, foco_n, titulo_foco(text, foco_n), descripcion),
                    )
                    focos += 1
                    continue
                if seccion is None:
                    continue
                descripcion = texto(text)
                if not descripcion:
                    continue
                orden[seccion] += 1
                estado, nota = norm_estado(row[9])
                gestion = texto(row[10])
                if nota:
                    gestion = nota if not gestion else f"{gestion}\n{nota}"
                cur = conn.execute(
                    """
                    INSERT INTO tareas (
                        semana_id, seccion, orden, descripcion, semana_prioridad,
                        fecha_entrega, estado, gestion, creado_en
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        semana_id,
                        seccion,
                        orden[seccion],
                        descripcion,
                        as_prioridad(row[6]),
                        as_fecha(row[7]),
                        estado,
                        gestion,
                        f"{inicio} 08:00:00",
                    ),
                )
                tarea_id = cur.lastrowid
                for nombre in separar_personas(row[8]):
                    conn.execute(
                        "INSERT OR IGNORE INTO tarea_responsable (tarea_id, personal_id) VALUES (?, ?)",
                        (tarea_id, id_persona(nombre)),
                    )
                tareas += 1
    finally:
        wb.close()

    ahora = datetime.now().replace(microsecond=0).isoformat(sep=" ")
    conn.execute(
        "INSERT INTO meta (clave, valor) VALUES ('origen', ?), ('importado_en', ?)",
        (os.path.basename(EXCEL), ahora),
    )
    return {"semanas": semanas, "tareas": tareas, "focos": focos, "personas": len(cache)}


def init_db():
    conn = connect()
    try:
        with open(SCHEMA, encoding="utf-8") as handle:
            conn.executescript(handle.read())
        conn.execute("PRAGMA foreign_keys = ON")
        total = conn.execute("SELECT COUNT(*) FROM semanas").fetchone()[0]
        if total == 0 and os.path.exists(EXCEL):
            resumen = importar_excel(conn)
            conn.commit()
            return (
                f"Excel cargado: {resumen['semanas']} semanas, "
                f"{resumen['tareas']} tareas, {resumen['focos']} focos, "
                f"{resumen['personas']} personas."
            )
        conn.commit()
        return f"Base lista: {total} semanas."
    finally:
        conn.close()


def dic(row):
    return dict(row) if row is not None else None


def tarea_json(db, tarea_id):
    tarea = db.execute("SELECT * FROM tareas WHERE id = ?", (tarea_id,)).fetchone()
    if not tarea:
        raise LookupError("Esa tarea no existe.")
    data = dict(tarea)
    data["responsables"] = [
        dict(item)
        for item in db.execute(
            """
            SELECT p.id, p.nombre
            FROM tarea_responsable tr
            JOIN personal p ON p.id = tr.personal_id
            WHERE tr.tarea_id = ?
            ORDER BY tr.rowid
            """,
            (tarea_id,),
        )
    ]
    return data


def tareas_de_semana(db, semana_id):
    filas = db.execute(
        """
        SELECT * FROM tareas
        WHERE semana_id = ?
        ORDER BY CASE seccion WHEN 'prioritaria' THEN 0 ELSE 1 END, orden, id
        """,
        (semana_id,),
    ).fetchall()
    if not filas:
        return []
    ids = [row["id"] for row in filas]
    marcas = ",".join("?" for _ in ids)
    responsables = {}
    for item in db.execute(
        f"""
        SELECT tr.tarea_id, p.id AS id, p.nombre
        FROM tarea_responsable tr
        JOIN personal p ON p.id = tr.personal_id
        WHERE tr.tarea_id IN ({marcas})
        ORDER BY tr.rowid
        """,
        ids,
    ):
        responsables.setdefault(item["tarea_id"], []).append(
            {"id": item["id"], "nombre": item["nombre"]}
        )
    salida = []
    for row in filas:
        data = dict(row)
        data["responsables"] = responsables.get(row["id"], [])
        salida.append(data)
    return salida


def resumen_tareas(tareas):
    realizadas = sum(1 for tarea in tareas if tarea["estado"] == "Realizado")
    minutos = sum(tarea["tiempo_minutos"] or 0 for tarea in tareas)
    return {
        "total": len(tareas),
        "realizadas": realizadas,
        "abiertas": len(tareas) - realizadas,
        "minutos": minutos,
    }


def exigir_estado(value):
    if value not in ESTADOS:
        raise ValueError("El estado no es válido.")
    return value


def exigir_seccion(value):
    if value not in SECCIONES:
        raise ValueError("La sección no es válida.")
    return value


def exigir_prioridad(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise ValueError("La semana de origen debe ser un número.")
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError("La semana de origen debe ser un número entero.")
        numero = int(value)
    elif isinstance(value, int):
        numero = value
    else:
        bruto = str(value).strip()
        if not re.fullmatch(r"\d{1,2}", bruto):
            raise ValueError("La semana de origen debe ser un número entero.")
        numero = int(bruto)
    if not 1 <= numero <= 53:
        raise ValueError("La semana de origen debe estar entre 1 y 53.")
    return numero


def exigir_minutos(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise ValueError("El tiempo no es válido.")
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError("El tiempo debe indicarse en minutos enteros.")
        minutos = int(value)
    else:
        try:
            minutos = int(value)
        except (TypeError, ValueError):
            raise ValueError("El tiempo debe indicarse en minutos enteros.")
    if minutos < 0 or minutos > 200 * 60:
        raise ValueError("El tiempo debe estar entre 0 y 200 horas.")
    return minutos


def exigir_texto(value, etiqueta, limite, obligatorio=False, minimo=0):
    cleaned = texto(value)
    if not cleaned:
        if obligatorio:
            raise ValueError(f"Falta completar {etiqueta[0].lower() + etiqueta[1:]}.")
        return None
    if minimo and len(cleaned) < minimo:
        raise ValueError(f"{etiqueta} debe tener al menos {minimo} caracteres.")
    if len(cleaned) > limite:
        raise ValueError(f"{etiqueta} supera el límite permitido.")
    return cleaned


def exigir_fecha(value):
    cleaned = exigir_texto(value, "La fecha de entrega", 80)
    if cleaned and re.fullmatch(r"\d{4}-\d{2}-\d{2}", cleaned):
        try:
            date.fromisoformat(cleaned)
        except ValueError:
            raise ValueError("La fecha de entrega no existe.")
    return cleaned


def reemplazar_responsables(db, tarea_id, ids):
    if not isinstance(ids, list):
        raise ValueError("Los responsables no tienen el formato esperado.")
    limpios = []
    for item in ids:
        try:
            numero = int(item)
        except (TypeError, ValueError):
            raise ValueError("Hay un responsable que no existe.")
        if numero not in limpios:
            limpios.append(numero)
    if limpios:
        marcas = ",".join("?" for _ in limpios)
        encontrados = db.execute(
            f"SELECT id FROM personal WHERE id IN ({marcas})",
            limpios,
        ).fetchall()
        if len(encontrados) != len(limpios):
            raise ValueError("Hay un responsable que no existe.")
    db.execute("DELETE FROM tarea_responsable WHERE tarea_id = ?", (tarea_id,))
    for persona_id in limpios:
        db.execute(
            "INSERT INTO tarea_responsable (tarea_id, personal_id) VALUES (?, ?)",
            (tarea_id, persona_id),
        )


def aplicar_tarea(db, tarea_id, data, crear=False):
    if crear or "descripcion" in data:
        descripcion = exigir_texto(data.get("descripcion"), "La descripción", 2000, True, minimo=3)
    else:
        descripcion = None
    seccion = exigir_seccion(data.get("seccion")) if (crear or "seccion" in data) else None
    estado = exigir_estado(data.get("estado", "Pendiente" if crear else None)) if (crear or "estado" in data) else None
    prioridad = exigir_prioridad(data.get("semana_prioridad")) if (crear or "semana_prioridad" in data) else None
    fecha = exigir_fecha(data.get("fecha_entrega")) if (crear or "fecha_entrega" in data) else None
    gestion = exigir_texto(data.get("gestion"), "La gestión", 4000) if (crear or "gestion" in data) else None
    minutos = exigir_minutos(data.get("tiempo_minutos")) if (crear or "tiempo_minutos" in data) else None
    return descripcion, seccion, estado, prioridad, fecha, gestion, minutos


@app.errorhandler(ValueError)
def error_valor(exc):
    return jsonify({"error": str(exc)}), 400


@app.errorhandler(LookupError)
def error_busqueda(exc):
    return jsonify({"error": str(exc)}), 404


@app.get("/")
def pagina():
    return render_template("index.html")


@app.get("/api/inicio")
def api_inicio():
    db = get_db()
    hoy = date.today()
    iso = hoy.isocalendar()
    semanas = []
    for row in db.execute(
        """
        SELECT s.*,
               COUNT(t.id) AS total,
               SUM(CASE WHEN t.estado = 'Realizado' THEN 1 ELSE 0 END) AS realizadas,
               SUM(COALESCE(t.tiempo_minutos, 0)) AS minutos
        FROM semanas s
        LEFT JOIN tareas t ON t.semana_id = s.id
        GROUP BY s.id
        ORDER BY s.anio DESC, s.numero DESC
        """
    ):
        item = dict(row)
        item["realizadas"] = item["realizadas"] or 0
        item["minutos"] = item["minutos"] or 0
        semanas.append(item)
    numeros = {item["numero"] for item in semanas if item["anio"] == ANIO_EXCEL}
    huecos = [n for n in range(1, (max(numeros) if numeros else 0) + 1) if n not in numeros]
    personal = [
        dict(row)
        for row in db.execute(
            """
            SELECT id, nombre, activo
            FROM personal
            ORDER BY CASE nombre
                WHEN 'Todos' THEN 2
                WHEN 'Desarrollo' THEN 1
                ELSE 0
            END, nombre COLLATE NOCASE
            """
        )
    ]
    meta = {row["clave"]: row["valor"] for row in db.execute("SELECT clave, valor FROM meta")}
    return jsonify(
        {
            "hoy": {"anio": iso.year, "semana": iso.week, "fecha": hoy.isoformat()},
            "semanas": semanas,
            "huecos": huecos,
            "personal": personal,
            "estados": list(ESTADOS),
            "meta": meta,
        }
    )


@app.get("/api/semanas/<int:semana_id>")
def api_semana(semana_id):
    db = get_db()
    semana = db.execute("SELECT * FROM semanas WHERE id = ?", (semana_id,)).fetchone()
    if not semana:
        raise LookupError("Esa semana no existe.")
    focos = [
        dict(row)
        for row in db.execute(
            "SELECT * FROM focos WHERE semana_id = ? ORDER BY orden, id",
            (semana_id,),
        )
    ]
    tareas = tareas_de_semana(db, semana_id)
    return jsonify(
        {
            "semana": dict(semana),
            "focos": focos,
            "tareas": tareas,
            "resumen": resumen_tareas(tareas),
        }
    )


@app.post("/api/semanas")
def api_crear_semana():
    db = get_db()
    try:
        origen = db.execute(
            "SELECT * FROM semanas ORDER BY anio DESC, numero DESC LIMIT 1"
        ).fetchone()
        if not origen:
            raise ValueError("No hay una semana desde la cual continuar.")
        anio, numero = siguiente_semana(origen["anio"], origen["numero"])
        if db.execute(
            "SELECT id FROM semanas WHERE anio = ? AND numero = ?",
            (anio, numero),
        ).fetchone():
            raise ValueError(f"La semana {numero} de {anio} ya existe.")
        inicio, fin = limites_semana(anio, numero)
        cur = db.execute(
            """
            INSERT INTO semanas (anio, numero, fecha_inicio, fecha_fin, es_copia)
            VALUES (?, ?, ?, ?, 0)
            """,
            (anio, numero, inicio, fin),
        )
        nueva_id = cur.lastrowid
        for foco in db.execute(
            "SELECT * FROM focos WHERE semana_id = ? ORDER BY orden",
            (origen["id"],),
        ):
            db.execute(
                """
                INSERT INTO focos (semana_id, orden, titulo, descripcion, estado)
                VALUES (?, ?, ?, ?, 'Sin estado')
                """,
                (nueva_id, foco["orden"], foco["titulo"], foco["descripcion"]),
            )
        copiadas = 0
        for tarea in db.execute(
            """
            SELECT * FROM tareas
            WHERE semana_id = ? AND estado != 'Realizado'
            ORDER BY seccion, orden
            """,
            (origen["id"],),
        ):
            cur = db.execute(
                """
                INSERT INTO tareas (
                    semana_id, seccion, orden, descripcion, semana_prioridad,
                    fecha_entrega, estado, gestion, tiempo_minutos
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    nueva_id,
                    tarea["seccion"],
                    tarea["orden"],
                    tarea["descripcion"],
                    tarea["semana_prioridad"],
                    tarea["fecha_entrega"],
                    tarea["estado"],
                    tarea["gestion"],
                    tarea["tiempo_minutos"],
                ),
            )
            nuevo_id = cur.lastrowid
            for rel in db.execute(
                "SELECT personal_id FROM tarea_responsable WHERE tarea_id = ? ORDER BY rowid",
                (tarea["id"],),
            ):
                db.execute(
                    "INSERT INTO tarea_responsable (tarea_id, personal_id) VALUES (?, ?)",
                    (nuevo_id, rel["personal_id"]),
                )
            copiadas += 1
        db.commit()
    except Exception:
        db.rollback()
        raise
    if copiadas == 1:
        mensaje = f"Semana {numero} de {anio} lista, con 1 tarea abierta."
    elif copiadas:
        mensaje = f"Semana {numero} de {anio} lista, con {copiadas} tareas abiertas."
    else:
        mensaje = f"Semana {numero} de {anio} lista. No había tareas abiertas para copiar."
    semana = dict(db.execute("SELECT * FROM semanas WHERE id = ?", (nueva_id,)).fetchone())
    return jsonify({"semana": semana, "copiadas": copiadas, "mensaje": mensaje})


@app.post("/api/semanas/<int:semana_id>/tareas")
def api_crear_tarea(semana_id):
    db = get_db()
    if not db.execute("SELECT id FROM semanas WHERE id = ?", (semana_id,)).fetchone():
        raise LookupError("Esa semana no existe.")
    data = request.get_json(silent=True) or {}
    try:
        descripcion, seccion, estado, prioridad, fecha, gestion, minutos = aplicar_tarea(
            db, None, data, crear=True
        )
        orden = db.execute(
            """
            SELECT COALESCE(MAX(orden), 0) + 1
            FROM tareas WHERE semana_id = ? AND seccion = ?
            """,
            (semana_id, seccion),
        ).fetchone()[0]
        cur = db.execute(
            """
            INSERT INTO tareas (
                semana_id, seccion, orden, descripcion, semana_prioridad,
                fecha_entrega, estado, gestion, tiempo_minutos
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (semana_id, seccion, orden, descripcion, prioridad, fecha, estado, gestion, minutos),
        )
        responsables = data.get("responsables") or []
        if not isinstance(responsables, list) or not responsables:
            raise ValueError("Asigna al menos una persona.")
        reemplazar_responsables(db, cur.lastrowid, responsables)
        db.commit()
        return jsonify({"tarea": tarea_json(db, cur.lastrowid)}), 201
    except Exception:
        db.rollback()
        raise


@app.patch("/api/tareas/<int:tarea_id>")
def api_editar_tarea(tarea_id):
    db = get_db()
    if not db.execute("SELECT id FROM tareas WHERE id = ?", (tarea_id,)).fetchone():
        raise LookupError("Esa tarea no existe.")
    data = request.get_json(silent=True) or {}
    try:
        campos = []
        params = []
        if "descripcion" in data:
            campos.append("descripcion = ?")
            params.append(exigir_texto(data.get("descripcion"), "La descripción", 2000, True, minimo=3))
        if "seccion" in data:
            campos.append("seccion = ?")
            params.append(exigir_seccion(data.get("seccion")))
        if "estado" in data:
            campos.append("estado = ?")
            params.append(exigir_estado(data.get("estado")))
        if "semana_prioridad" in data:
            campos.append("semana_prioridad = ?")
            params.append(exigir_prioridad(data.get("semana_prioridad")))
        if "fecha_entrega" in data:
            campos.append("fecha_entrega = ?")
            params.append(exigir_fecha(data.get("fecha_entrega")))
        if "gestion" in data:
            campos.append("gestion = ?")
            params.append(exigir_texto(data.get("gestion"), "La gestión", 4000))
        if "tiempo_minutos" in data:
            campos.append("tiempo_minutos = ?")
            params.append(exigir_minutos(data.get("tiempo_minutos")))
        if not campos and "responsables" not in data:
            raise ValueError("No hay cambios para guardar.")
        if campos:
            params.append(tarea_id)
            db.execute(f"UPDATE tareas SET {', '.join(campos)} WHERE id = ?", params)
        if "responsables" in data:
            reemplazar_responsables(db, tarea_id, data.get("responsables") or [])
        db.commit()
        return jsonify({"tarea": tarea_json(db, tarea_id)})
    except Exception:
        db.rollback()
        raise


@app.delete("/api/tareas/<int:tarea_id>")
def api_eliminar_tarea(tarea_id):
    db = get_db()
    cur = db.execute("DELETE FROM tareas WHERE id = ?", (tarea_id,))
    db.commit()
    if cur.rowcount == 0:
        raise LookupError("Esa tarea no existe.")
    return jsonify({"ok": True})


@app.post("/api/semanas/<int:semana_id>/focos")
def api_crear_foco(semana_id):
    db = get_db()
    if not db.execute("SELECT id FROM semanas WHERE id = ?", (semana_id,)).fetchone():
        raise LookupError("Esa semana no existe.")
    data = request.get_json(silent=True) or {}
    try:
        total = db.execute(
            "SELECT COUNT(*) FROM focos WHERE semana_id = ?",
            (semana_id,),
        ).fetchone()[0]
        if total >= 6:
            raise ValueError("Esta semana ya tiene el máximo de actividades de cabecera.")
        orden = total + 1
        titulo = exigir_texto(data.get("titulo") or f"Actividad {orden}", "El título", 80, True, minimo=2)
        descripcion = exigir_texto(data.get("descripcion"), "La descripción", 500, True, minimo=3)
        estado = exigir_estado(data.get("estado") or "Sin estado")
        cur = db.execute(
            """
            INSERT INTO focos (semana_id, orden, titulo, descripcion, estado, detalle)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                semana_id,
                orden,
                titulo,
                descripcion,
                estado,
                exigir_texto(data.get("detalle"), "El detalle", 1000),
            ),
        )
        db.commit()
        foco = dict(db.execute("SELECT * FROM focos WHERE id = ?", (cur.lastrowid,)).fetchone())
        return jsonify({"foco": foco}), 201
    except Exception:
        db.rollback()
        raise


@app.patch("/api/focos/<int:foco_id>")
def api_editar_foco(foco_id):
    db = get_db()
    if not db.execute("SELECT id FROM focos WHERE id = ?", (foco_id,)).fetchone():
        raise LookupError("Esa actividad no existe.")
    data = request.get_json(silent=True) or {}
    try:
        campos = []
        params = []
        if "titulo" in data:
            campos.append("titulo = ?")
            params.append(exigir_texto(data.get("titulo"), "El título", 80, True, minimo=2))
        if "descripcion" in data:
            campos.append("descripcion = ?")
            params.append(exigir_texto(data.get("descripcion"), "La descripción", 500, True, minimo=3))
        if "estado" in data:
            campos.append("estado = ?")
            params.append(exigir_estado(data.get("estado")))
        if "detalle" in data:
            campos.append("detalle = ?")
            params.append(exigir_texto(data.get("detalle"), "El detalle", 1000))
        if not campos:
            raise ValueError("No hay cambios para guardar.")
        params.append(foco_id)
        db.execute(f"UPDATE focos SET {', '.join(campos)} WHERE id = ?", params)
        db.commit()
        foco = dict(db.execute("SELECT * FROM focos WHERE id = ?", (foco_id,)).fetchone())
        return jsonify({"foco": foco})
    except Exception:
        db.rollback()
        raise


@app.delete("/api/focos/<int:foco_id>")
def api_eliminar_foco(foco_id):
    db = get_db()
    cur = db.execute("DELETE FROM focos WHERE id = ?", (foco_id,))
    db.commit()
    if cur.rowcount == 0:
        raise LookupError("Esa actividad no existe.")
    return jsonify({"ok": True})


@app.get("/api/buscar")
def api_buscar():
    consulta = (request.args.get("q") or "").strip()
    if len(fold(consulta)) < 2:
        return jsonify({"resultados": []})
    needle = fold(consulta)
    db = get_db()
    resultados = []
    filas = db.execute(
        """
        SELECT t.id, t.descripcion, t.gestion, t.estado, t.seccion,
               s.id AS semana_id, s.numero, s.anio,
               (
                   SELECT GROUP_CONCAT(p.nombre, ' ')
                   FROM tarea_responsable tr
                   JOIN personal p ON p.id = tr.personal_id
                   WHERE tr.tarea_id = t.id
               ) AS responsables
        FROM tareas t
        JOIN semanas s ON s.id = t.semana_id
        ORDER BY s.anio DESC, s.numero DESC, t.orden
        """
    )
    for row in filas:
        blob = " ".join(
            part for part in (row["descripcion"], row["gestion"], row["responsables"]) if part
        )
        if needle not in fold(blob):
            continue
        resultados.append(
            {
                "id": row["id"],
                "descripcion": row["descripcion"],
                "estado": row["estado"],
                "seccion": row["seccion"],
                "semana_id": row["semana_id"],
                "numero": row["numero"],
                "anio": row["anio"],
            }
        )
        if len(resultados) >= 30:
            break
    return jsonify({"resultados": resultados})


@app.post("/api/personal")
def api_crear_personal():
    db = get_db()
    data = request.get_json(silent=True) or {}
    try:
        persona = obtener_o_crear_persona(db, data.get("nombre"))
        db.commit()
        return jsonify({"persona": dict(persona)}), 201
    except sqlite3.IntegrityError:
        db.rollback()
        raise ValueError("Esa persona ya existe.")
    except Exception:
        db.rollback()
        raise


@app.patch("/api/personal/<int:persona_id>")
def api_editar_personal(persona_id):
    db = get_db()
    if not db.execute("SELECT id FROM personal WHERE id = ?", (persona_id,)).fetchone():
        raise LookupError("Esa persona no existe.")
    data = request.get_json(silent=True) or {}
    if "activo" not in data:
        raise ValueError("No hay cambios para guardar.")
    activo = 1 if data.get("activo") in (1, True, "1") else 0
    try:
        db.execute("UPDATE personal SET activo = ? WHERE id = ?", (activo, persona_id))
        db.commit()
        persona = dict(db.execute("SELECT * FROM personal WHERE id = ?", (persona_id,)).fetchone())
        return jsonify({"persona": persona})
    except Exception:
        db.rollback()
        raise


def main():
    if "--reimport" in sys.argv and os.path.exists(DB_PATH):
        os.remove(DB_PATH)
    print(init_db())
    if "--servidor" in sys.argv:
        host = "0.0.0.0"
        port = 5099
        url = "http://192.168.20.205:5099"
    else:
        host = "127.0.0.1"
        port = int(os.environ.get("PORT", "5050"))
        url = f"http://127.0.0.1:{port}"
    print(f"Plan semanal listo en {url}")
    if "--no-browser" not in sys.argv:
        threading.Thread(
            target=lambda: (time.sleep(1.2), webbrowser.open(url)),
            daemon=True,
        ).start()
    app.run(host=host, port=port, debug=False, threaded=True)


if __name__ == "__main__":
    main()
